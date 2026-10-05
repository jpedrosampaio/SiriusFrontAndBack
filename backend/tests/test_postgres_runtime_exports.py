import os
import unittest
from datetime import timedelta
from io import BytesIO
from uuid import UUID
from unittest.mock import patch
from openpyxl import load_workbook
from db.models.health import Meal
from db.models.studies import StudySession
from db.session import unit_of_work
import test_postgres_runtime_nutrition as setup
import test_postgres_runtime_catalog as catalog


@unittest.skipUnless(os.getenv('RUN_POSTGRES_TESTS')=='true','Disposable PostgreSQL required')
class RuntimeExports(unittest.IsolatedAsyncioTestCase):
    ok=setup.RuntimeNutrition.ok
    asyncTearDown=setup.RuntimeNutrition.asyncTearDown

    async def asyncSetUp(self):
        await setup.RuntimeNutrition.asyncSetUp(self)
        async def account(request):return {'user_id':str(self.bob if request.headers.get('Authorization')=='Bearer bob' else self.uid),'timezone':'America/Sao_Paulo','name':'A <B> & C'}
        p=patch('services.domain_exports.account',account);p.start();self.patches.append(p)

    async def test_nutrition_full_export_owner_period_and_empty_pdf(self):
        async with unit_of_work() as session:
            session.add_all([Meal(user_id=self.uid,date=self.today,name='Meal',meal_type='=1+1',reported_calories=350,reported_protein=25) for _ in range(1005)])
            session.add(Meal(user_id=self.uid,date=self.today-timedelta(days=1),name='Old',meal_type='old'))
            session.add(Meal(user_id=self.bob,date=self.today,name='Foreign',meal_type='foreign'))
        query={'start':self.today.isoformat(),'end':self.today.isoformat()}
        response=await self.http.get('/api/export/nutrition/excel',params=query);self.assertEqual(response.status_code,200,response.text)
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row,1006);self.assertEqual(sheet['C2'].value,350);self.assertEqual(sheet['D2'].value,25)
        self.assertEqual(sheet['B2'].data_type,'s')
        foreign=await self.http.get('/api/export/nutrition/excel',params=query,headers={'Authorization':'Bearer bob'})
        self.assertEqual(load_workbook(BytesIO(foreign.content)).active.max_row,2)
        future=(self.today+timedelta(days=1)).isoformat()
        pdf=await self.http.get('/api/export/nutrition/pdf',params={'start':future});self.assertEqual(pdf.status_code,200);self.assertTrue(pdf.content.startswith(b'%PDF'))
        self.assertEqual((await self.http.get('/api/export/nutrition/excel',params={'start':future,'end':self.today.isoformat()})).status_code,422)

    async def test_study_export_sessions_counts_and_pdf(self):
        _,_,book=await catalog.RuntimeCatalog.setup_catalog(self)
        async with unit_of_work() as session:
            session.add_all([StudySession(user_id=self.uid,notebook_id=UUID(book['notebook_id']),date=self.today,duration_minutes=25,completed=True),
                StudySession(user_id=self.uid,notebook_id=UUID(book['notebook_id']),date=self.today-timedelta(days=1),duration_minutes=50,completed=True),
                StudySession(user_id=self.uid,notebook_id=UUID(book['notebook_id']),date=self.today,duration_minutes=10,completed=False),
                StudySession(user_id=self.bob,date=self.today,duration_minutes=100,completed=True)])
        query={'start':self.today.isoformat(),'end':self.today.isoformat()}
        response=await self.http.get('/api/export/study/excel',params=query);self.assertEqual(response.status_code,200,response.text)
        wb=load_workbook(BytesIO(response.content));sheet=wb.worksheets[0]
        self.assertEqual(sheet['D2'].value,25);self.assertEqual(sheet['D4'].value,25)
        self.assertEqual(wb.worksheets[1]['D2'].value,1)
        pdf=await self.http.get('/api/export/study/pdf',params=query);self.assertEqual(pdf.status_code,200,pdf.text[:100]);self.assertTrue(pdf.content.startswith(b'%PDF'))
