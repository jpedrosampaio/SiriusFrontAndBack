"""Study and nutrition exports over SQL-filtered facts."""
from datetime import date as Date,datetime
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID
from xml.sax.saxutils import escape
from fastapi import APIRouter,Request,HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select,func
from sqlalchemy.orm import selectinload
from db.models.health import Meal
from db.models.studies import Notebook,StudyArea,StudySession,Flashcard
from db.session import unit_of_work
from services.auth_routes import account
from services.nutrition import meal_json

router=APIRouter()


def date_filter(query,column,start,end):
    if start:query=query.where(column>=start)
    if end:query=query.where(column<=end)
    return query


async def nutrition_data(uid,start,end):
    async with unit_of_work() as session:
        query=date_filter(select(Meal).options(selectinload(Meal.items)).where(Meal.user_id==uid),Meal.date,start,end)
        return [meal_json(row) for row in (await session.scalars(query.order_by(Meal.date.desc(),Meal.id))).all()]


async def study_data(uid,start,end):
    async with unit_of_work() as session:
        query=date_filter(select(StudySession,Notebook.name).outerjoin(Notebook,(Notebook.id==StudySession.notebook_id)&(Notebook.user_id==StudySession.user_id))
            .where(StudySession.user_id==uid,StudySession.completed.is_(True)),StudySession.date,start,end)
        sessions=[{'date':row.date.isoformat(),'notebook_id':str(row.notebook_id),'notebook_name':name or '',
            'session_type':row.source,'duration_minutes':row.duration_minutes,'pomodoros':1 if row.source=='focus' else 0}
            for row,name in (await session.execute(query.order_by(StudySession.date.desc(),StudySession.id))).all()]
        counts_query=date_filter(select(StudySession.notebook_id,func.count()).where(StudySession.user_id==uid,StudySession.completed.is_(True)),StudySession.date,start,end)
        counts=dict((await session.execute(counts_query.group_by(StudySession.notebook_id))).all())
        cards=dict((await session.execute(select(Flashcard.notebook_id,func.count()).join(Notebook,(Notebook.id==Flashcard.notebook_id)&(Notebook.user_id==Flashcard.user_id))
            .where(Flashcard.user_id==uid,Flashcard.archived_at.is_(None),Notebook.archived_at.is_(None)).group_by(Flashcard.notebook_id))).all())
        books=[{'notebook_id':str(row.id),'name':row.name,'area_name':area,'flashcards_count':cards.get(row.id,0),'sessions_count':counts.get(row.id,0)}
            for row,area in (await session.execute(select(Notebook,StudyArea.name).join(StudyArea,(StudyArea.id==Notebook.area_id)&(StudyArea.user_id==Notebook.user_id))
                .where(Notebook.user_id==uid,Notebook.archived_at.is_(None)).order_by(Notebook.name,Notebook.id))).all()]
        return books,sessions,sum(cards.values())


@router.get('/export/study/{format}')
async def export_study(request: Request, format: str, start: Date | None=None, end: Date | None=None):
    """Export study data as PDF or Excel"""
    profile = await account(request)
    user = SimpleNamespace(**profile)
    if format not in ('excel', 'pdf'): raise HTTPException(400, 'Formato deve ser excel ou pdf')
    if start and end and end < start: raise HTTPException(422, 'Periodo invalido')
    notebooks, sessions, flashcards_count = await study_data(UUID(user.user_id), start, end)
    if format == 'excel':
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = 'Sessões de Estudo'
        header_fill = PatternFill(start_color='A855F7', end_color='A855F7', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        headers = ['Data', 'Matéria', 'Tipo', 'Duração (min)', 'Pomodoros']
        for col, h in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        for row, s in enumerate(sessions, 2):
            ws1.cell(row=row, column=1, value=s.get('date', '')).border = thin_border
            ws1.cell(row=row, column=2, value=s.get('notebook_name', '')).border = thin_border
            ws1.cell(row=row, column=3, value=s.get('session_type', 'study')).border = thin_border
            ws1.cell(row=row, column=4, value=s.get('duration_minutes', 0)).border = thin_border
            ws1.cell(row=row, column=5, value=s.get('pomodoros', 0)).border = thin_border
        total_row = len(sessions) + 3
        ws1.cell(row=total_row, column=3, value='TOTAL:').font = Font(bold=True)
        ws1.cell(row=total_row, column=4, value=sum((s.get('duration_minutes', 0) for s in sessions))).font = Font(bold=True)
        ws2 = wb.create_sheet('Matérias')
        headers2 = ['Matéria', 'Área', 'Flashcards', 'Sessões']
        for col, h in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        for row, nb in enumerate(notebooks, 2):
            nb_id = nb.get('notebook_id', '')
            nb_flashcards = nb['flashcards_count']
            nb_sessions = nb['sessions_count']
            ws2.cell(row=row, column=1, value=nb.get('name', '')).border = thin_border
            ws2.cell(row=row, column=2, value=nb.get('area_name', '')).border = thin_border
            ws2.cell(row=row, column=3, value=nb_flashcards).border = thin_border
            ws2.cell(row=row, column=4, value=nb_sessions).border = thin_border
        for ws in [ws1, ws2]:
            for col in ws.columns:
                max_len = max((len(str(cell.value or '')) for cell in col))
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)
        buf = BytesIO()
        for sheet in wb:
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str): cell.data_type = 's'
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=estudos_sirius.xlsx'})
    elif format == 'pdf':
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        styles = getSampleStyleSheet()
        elements = []
        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=18, textColor=colors.HexColor('#A855F7'))
        elements.append(Paragraph('Relatório de Estudos - Sirius', title_style))
        elements.append(Spacer(1, 12))
        elements.append(Paragraph(f'Usuário: {escape(user.name)}', styles['Normal']))
        elements.append(Paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y')}", styles['Normal']))
        elements.append(Spacer(1, 20))
        elements.append(Paragraph('Resumo por Matéria', styles['Heading2']))
        nb_data = [['Matéria', 'Área', 'Flashcards', 'Sessões']]
        for nb in notebooks:
            nb_id = nb.get('notebook_id', '')
            nb_data.append([nb.get('name', '')[:25], nb.get('area_name', '')[:20], str(nb['flashcards_count']), str(nb['sessions_count'])])
        table = Table(nb_data, repeatRows=1, colWidths=[5 * cm, 4 * cm, 3 * cm, 3 * cm])
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#A855F7')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')])]))
        elements.append(table)
        elements.append(Spacer(1, 20))
        total_minutes = sum((s.get('duration_minutes', 0) for s in sessions))
        total_hours = total_minutes / 60
        elements.append(Paragraph(f'Total de horas estudadas: {total_hours:.1f}h ({total_minutes} min)', styles['Normal']))
        elements.append(Paragraph(f'Total de sessões: {len(sessions)}', styles['Normal']))
        elements.append(Paragraph(f'Total de flashcards: {flashcards_count}', styles['Normal']))
        doc.build(elements)
        buf.seek(0)
        return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=estudos_sirius.pdf'})
    raise HTTPException(status_code=400, detail="Formato deve ser 'excel' ou 'pdf'")

@router.get('/export/nutrition/{format}')
async def export_nutrition(request: Request, format: str, start: Date | None=None, end: Date | None=None):
    """Export nutrition data as PDF or Excel"""
    profile = await account(request)
    user = SimpleNamespace(**profile)
    if format not in ('excel', 'pdf'): raise HTTPException(400, 'Formato deve ser excel ou pdf')
    if start and end and end < start: raise HTTPException(422, 'Periodo invalido')
    meals = await nutrition_data(UUID(user.user_id), start, end)
    if format == 'excel':
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Refeições'
        header_fill = PatternFill(start_color='22C55E', end_color='22C55E', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        headers = ['Data', 'Refeição', 'Calorias', 'Proteína(g)', 'Carbos(g)', 'Gordura(g)']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
        for row, m in enumerate(meals, 2):
            ws.cell(row=row, column=1, value=m.get('date', '')).border = thin_border
            ws.cell(row=row, column=2, value=m.get('meal_type', '')).border = thin_border
            ws.cell(row=row, column=3, value=round(m.get('total_calories', 0), 1)).border = thin_border
            ws.cell(row=row, column=4, value=round(m.get('total_protein', 0), 1)).border = thin_border
            ws.cell(row=row, column=5, value=round(m.get('total_carbs', 0), 1)).border = thin_border
            ws.cell(row=row, column=6, value=round(m.get('total_fat', 0), 1)).border = thin_border
        for col in ws.columns:
            max_len = max((len(str(cell.value or '')) for cell in col))
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 25)
        buf = BytesIO()
        for sheet in wb:
            for cells in sheet:
                for cell in cells:
                    if isinstance(cell.value, str): cell.data_type = 's'
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=nutricao_sirius.xlsx'})
    elif format == 'pdf':
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        styles = getSampleStyleSheet()
        elements = []
        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=18, textColor=colors.HexColor('#22C55E'))
        elements.append(Paragraph('Relatório Nutricional - Sirius', title_style))
        elements.append(Spacer(1, 12))
        elements.append(Paragraph(f'Usuário: {escape(user.name)}', styles['Normal']))
        elements.append(Spacer(1, 20))
        data = [['Data', 'Refeição', 'Calorias', 'Proteína', 'Carbos', 'Gordura']]
        for m in meals:
            data.append([m.get('date', ''), m.get('meal_type', ''), f"{m.get('total_calories', 0):.0f}", f"{m.get('total_protein', 0):.1f}g", f"{m.get('total_carbs', 0):.1f}g", f"{m.get('total_fat', 0):.1f}g"])
        table = Table(data, repeatRows=1, colWidths=[2.5 * cm, 3 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm, 2.5 * cm])
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#22C55E')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 8), ('GRID', (0, 0), (-1, -1), 0.5, colors.grey), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')])]))
        elements.append(table)
        doc.build(elements)
        buf.seek(0)
        return StreamingResponse(buf, media_type='application/pdf', headers={'Content-Disposition': 'attachment; filename=nutricao_sirius.pdf'})
    raise HTTPException(status_code=400, detail="Formato deve ser 'excel' ou 'pdf'")
