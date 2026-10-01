from datetime import date as Date, datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID
from xml.sax.saxutils import escape
from fastapi import APIRouter, Request, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from db.models.finance import FinancialTransaction
from db.session import unit_of_work
from services.auth_routes import account

router = APIRouter()

@router.get('/export/finance/{format}')
async def export_finance(request: Request, format: str, start: Date | None = None, end: Date | None = None):
    """Export financial data as PDF or Excel"""
    profile = await account(request)
    if start and end and end < start: raise HTTPException(422,'Período inválido.')
    user = SimpleNamespace(**profile)
    query = select(FinancialTransaction).where(FinancialTransaction.user_id == UUID(user.user_id))
    if start: query = query.where(FinancialTransaction.date >= start)
    if end: query = query.where(FinancialTransaction.date <= end)
    async with unit_of_work() as session:
        rows = (await session.scalars(query.order_by(FinancialTransaction.date.desc(),FinancialTransaction.id))).all()
        transactions = [{'date':r.date.isoformat(),'description':r.description or '', 'category':r.category,
                         'type':r.type,'amount':r.amount} for r in rows]

    if format == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Transações"
        
        # Header styling
        header_fill = PatternFill(start_color="007AFF", end_color="007AFF", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin")
        )
        
        headers = ["Data", "Descrição", "Categoria", "Tipo", "Valor (R$)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
        
        total_income = 0
        total_expenses = 0
        for row, t in enumerate(transactions, 2):
            ws.cell(row=row, column=1, value=t.get("date", "")).border = thin_border
            ws.cell(row=row, column=2, value=t.get("description", "")).border = thin_border
            ws.cell(row=row, column=3, value=t.get("category", "")).border = thin_border
            tipo = "Receita" if t.get("type") == "income" else "Despesa"
            ws.cell(row=row, column=4, value=tipo).border = thin_border
            amount = t.get("amount", 0)
            ws.cell(row=row, column=5, value=amount).border = thin_border
            ws.cell(row=row, column=5).number_format = '#,##0.00'
            if t.get("type") == "income":
                total_income += amount
            else:
                total_expenses += amount
        
        # Summary row
        summary_row = len(transactions) + 3
        ws.cell(row=summary_row, column=3, value="TOTAL RECEITAS:").font = Font(bold=True)
        ws.cell(row=summary_row, column=5, value=total_income).font = Font(bold=True, color="00AA00")
        ws.cell(row=summary_row, column=5).number_format = '#,##0.00'
        ws.cell(row=summary_row + 1, column=3, value="TOTAL DESPESAS:").font = Font(bold=True)
        ws.cell(row=summary_row + 1, column=5, value=total_expenses).font = Font(bold=True, color="FF0000")
        ws.cell(row=summary_row + 1, column=5).number_format = '#,##0.00'
        ws.cell(row=summary_row + 2, column=3, value="SALDO:").font = Font(bold=True)
        ws.cell(row=summary_row + 2, column=5, value=total_income - total_expenses).font = Font(bold=True)
        ws.cell(row=summary_row + 2, column=5).number_format = '#,##0.00'
        
        # Auto-fit columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 30)
        
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=financas_sirius.xlsx"}
        )
    
    elif format == "pdf":
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        
        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        elements = []
        
        title_style = ParagraphStyle('Title', parent=styles['Title'], fontSize=18, textColor=colors.HexColor('#007AFF'))
        elements.append(Paragraph("Relatório Financeiro - Sirius", title_style))
        elements.append(Spacer(1, 12))
        elements.append(Paragraph(f"Usuário: {escape(user.name)}", styles['Normal']))
        elements.append(Paragraph(f"Data: {datetime.now().strftime('%d/%m/%Y')}", styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Table
        data = [["Data", "Descrição", "Categoria", "Tipo", "Valor"]]
        total_income = 0
        total_expenses = 0
        for t in transactions:
            tipo = "Receita" if t.get("type") == "income" else "Despesa"
            amount = t.get("amount", 0)
            data.append([
                t.get("date", ""),
                t.get("description", "")[:30],
                t.get("category", ""),
                tipo,
                f"R$ {amount:.2f}"
            ])
            if t.get("type") == "income":
                total_income += amount
            else:
                total_expenses += amount
        
        data.append(["", "", "", "RECEITAS:", f"R$ {total_income:.2f}"])
        data.append(["", "", "", "DESPESAS:", f"R$ {total_expenses:.2f}"])
        data.append(["", "", "", "SALDO:", f"R$ {total_income - total_expenses:.2f}"])
        
        table = Table(data, colWidths=[2.5*cm, 5*cm, 3*cm, 2.5*cm, 3*cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#007AFF')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -4), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -4), [colors.white, colors.HexColor('#F5F5F5')]),
            ('FONTNAME', (3, -3), (-1, -1), 'Helvetica-Bold'),
        ]))
        elements.append(table)
        doc.build(elements)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=financas_sirius.pdf"}
        )
    
    raise HTTPException(status_code=400, detail="Formato deve ser 'excel' ou 'pdf'")

